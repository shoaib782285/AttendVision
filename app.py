import os
import sqlite3
import base64
import re
from io import BytesIO
from datetime import datetime
from pathlib import Path
from functools import wraps

import cv2
import numpy as np

from flask import (
    Flask, render_template, request,
    redirect, url_for, session, flash, jsonify, send_file
)
from werkzeug.security import generate_password_hash, check_password_hash
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment


# ==================================================
# 1. APP CONFIGURATION
# ==================================================

app = Flask(__name__)

app.secret_key = os.environ.get(
    "ATTENDANCE_SECRET_KEY",
    os.urandom(32)
)

app.config["MAX_CONTENT_LENGTH"] = 8 * 1024 * 1024

DB_PATH = os.path.join(app.root_path, "attendance.db")

FACE_DIR = Path(app.root_path) / "face_data"
FACE_DIR.mkdir(parents=True, exist_ok=True)

FACE_MODEL_PATH = Path(app.root_path) / "face_model.yml"
FACE_MODEL_META_PATH = Path(app.root_path) / "face_model.version"
FACE_MODEL_VERSION = "3"
FACE_CONFIDENCE_THRESHOLD = float(
    os.environ.get("FACE_CONFIDENCE_THRESHOLD", "110")
)
COLLEGE_START = datetime.strptime(
    os.environ.get("COLLEGE_START", "08:30"),
    "%H:%M",
).time()
COLLEGE_END = datetime.strptime(
    os.environ.get("COLLEGE_END", "16:30"),
    "%H:%M",
).time()
LATE_AFTER = datetime.strptime(
    os.environ.get("LATE_AFTER", "09:00"),
    "%H:%M",
).time()

ADMIN_USERNAME = os.environ.get("ADMIN_USERNAME", "admin")
ADMIN_PASSWORD = os.environ.get("ADMIN_PASSWORD", "Admin@123")
ADMIN_PASSWORD_HASH = generate_password_hash(ADMIN_PASSWORD)

FACE_CASCADE_PATH = (
    cv2.data.haarcascades
    + "haarcascade_frontalface_default.xml"
)

face_detector = cv2.CascadeClassifier(FACE_CASCADE_PATH)


def face_model_is_current():
    if not FACE_MODEL_PATH.is_file() or not FACE_MODEL_META_PATH.is_file():
        return False
    if FACE_MODEL_META_PATH.read_text(encoding="ascii").strip() != FACE_MODEL_VERSION:
        return False
    face_files = [
        path for path in FACE_DIR.glob("*.jpg") if path.is_file()
    ]
    return bool(face_files) and all(
        FACE_MODEL_PATH.stat().st_mtime >= path.stat().st_mtime
        for path in face_files
    )


# ==================================================
# 2. DATABASE
# ==================================================

def get_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    return conn


def init_db():
    with get_db() as conn:

        # Existing student records are preserved.
        conn.execute("""
            CREATE TABLE IF NOT EXISTS students (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT NOT NULL,
                roll_number TEXT NOT NULL UNIQUE,
                email TEXT,
                branch TEXT,
                academic_year TEXT,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )
        """)

        conn.execute("""
            CREATE TABLE IF NOT EXISTS attendance (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                student_id INTEGER NOT NULL,
                attendance_date TEXT NOT NULL,
                attendance_time TEXT NOT NULL,
                entry_time TEXT,
                exit_time TEXT,
                status TEXT NOT NULL DEFAULT 'Present',
                UNIQUE(student_id, attendance_date),
                FOREIGN KEY (student_id) REFERENCES students(id)
            )
        """)
        student_columns = {
            row["name"]
            for row in conn.execute("PRAGMA table_info(students)").fetchall()
        }
        if "branch" not in student_columns:
            conn.execute("ALTER TABLE students ADD COLUMN branch TEXT")
        if "academic_year" not in student_columns:
            conn.execute("ALTER TABLE students ADD COLUMN academic_year TEXT")

        attendance_columns = {
            row["name"]
            for row in conn.execute("PRAGMA table_info(attendance)").fetchall()
        }
        if "entry_time" not in attendance_columns:
            conn.execute("ALTER TABLE attendance ADD COLUMN entry_time TEXT")
        if "exit_time" not in attendance_columns:
            conn.execute("ALTER TABLE attendance ADD COLUMN exit_time TEXT")
        conn.execute(
            """
            UPDATE attendance
            SET entry_time = COALESCE(entry_time, attendance_time)
            WHERE entry_time IS NULL
            """
        )

init_db()


# ==================================================
# 3. LOGIN PROTECTION
# ==================================================

def login_required(view):
    @wraps(view)
    def wrapped_view(*args, **kwargs):
        if "admin" not in session:
            session.clear()
            return redirect(url_for("login"))

        return view(*args, **kwargs)

    return wrapped_view


# ==================================================
# 4. LOGIN
# ==================================================

@app.route("/", methods=["GET", "POST"])
def login():

    if request.method == "GET" and "admin" in session:
        return redirect(url_for("dashboard"))

    error = None

    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "")

        if (
            username == ADMIN_USERNAME
            and check_password_hash(ADMIN_PASSWORD_HASH, password)
        ):
            session.clear()
            session["admin"] = username
            return redirect(url_for("dashboard"))

        error = "Invalid username or password."

    return render_template("login.html", error=error)


# ==================================================
# 5. DASHBOARD
# ==================================================

@app.route("/dashboard")
@login_required
def dashboard():

    with get_db() as conn:
        students = conn.execute(
            "SELECT * FROM students ORDER BY id DESC"
        ).fetchall()

        today = datetime.now().strftime("%Y-%m-%d")

        today_count = conn.execute(
            """
            SELECT COUNT(*) AS total
            FROM attendance
            WHERE attendance_date = ?
            """,
            (today,)
        ).fetchone()["total"]

        week_start = datetime.now().date().fromordinal(
            datetime.now().date().toordinal() - 6
        ).isoformat()
        week_rows = conn.execute(
            """
            SELECT attendance_date, COUNT(*) AS total
            FROM attendance
            WHERE attendance_date >= ?
            GROUP BY attendance_date
            ORDER BY attendance_date
            """,
            (week_start,),
        ).fetchall()

        recent_records = conn.execute(
            """
            SELECT students.name, students.roll_number,
                   students.branch, students.academic_year,
                   attendance.attendance_date,
                   attendance.entry_time,
                   attendance.exit_time,
                   attendance.status
            FROM attendance
            JOIN students ON students.id = attendance.student_id
            ORDER BY attendance.id DESC
            LIMIT 6
            """
        ).fetchall()

    week_counts = {
        row["attendance_date"]: row["total"]
        for row in week_rows
    }
    attendance_chart = []
    for offset in range(6, -1, -1):
        day = datetime.now().date()
        chart_date = day.fromordinal(day.toordinal() - offset)
        date_key = chart_date.isoformat()
        attendance_chart.append({
            "label": chart_date.strftime("%a"),
            "date": date_key,
            "count": week_counts.get(date_key, 0),
        })

    face_ids = {
        path.stem
        for path in FACE_DIR.glob("*.jpg")
        if path.is_file()
    }

    return render_template(
        "dashboard.html",
        username=session["admin"],
        students=students,
        total_students=len(students),
        today_count=today_count,
        face_ids=face_ids,
        model_ready=face_model_is_current(),
        face_count=len(face_ids),
        attendance_window="08:30 AM - 04:30 PM",
        attendance_chart=attendance_chart,
        recent_records=recent_records,
    )


@app.route("/api/dashboard-stats")
@login_required
def dashboard_stats():
    """Return lightweight dashboard metrics for live refreshes."""
    today = datetime.now().strftime("%Y-%m-%d")
    with get_db() as conn:
        total_students = conn.execute(
            "SELECT COUNT(*) AS total FROM students"
        ).fetchone()["total"]
        today_count = conn.execute(
            """
            SELECT COUNT(*) AS total
            FROM attendance
            WHERE attendance_date = ?
            """,
            (today,)
        ).fetchone()["total"]
        total_attendance = conn.execute(
            "SELECT COUNT(*) AS total FROM attendance"
        ).fetchone()["total"]

    return jsonify({
        "success": True,
        "total_students": total_students,
        "today_count": today_count,
        "total_attendance": total_attendance,
        "face_count": len([
            path for path in FACE_DIR.glob("*.jpg") if path.is_file()
        ]),
        "model_ready": face_model_is_current(),
    })


# ==================================================
# 6. REGISTER STUDENT
# ==================================================

@app.route("/students/add", methods=["POST"])
@login_required
def add_student():

    name = request.form.get("name", "").strip()
    roll_number = request.form.get("roll_number", "").strip()
    email = request.form.get("email", "").strip()
    branch = request.form.get("branch", "").strip()
    academic_year = request.form.get("academic_year", "").strip()

    if not name or not roll_number:
        flash("Student name and roll number are required.", "error")
        return redirect(url_for("dashboard"))

    try:
        with get_db() as conn:
            conn.execute(
                """
                INSERT INTO students
                    (name, roll_number, email, branch, academic_year)
                VALUES (?, ?, ?, ?, ?)
                """,
                (
                    name,
                    roll_number,
                    email or None,
                    branch or None,
                    academic_year or None,
                )
            )

        flash("Student successfully registered.", "success")

    except sqlite3.IntegrityError:
        flash("That roll number is already registered.", "error")

    return redirect(url_for("dashboard"))


# ==================================================
# 7. DELETE STUDENT
# ==================================================

@app.route("/students/<int:student_id>/delete", methods=["POST"])
@login_required
def delete_student(student_id):
    face_path = FACE_DIR / f"{student_id}.jpg"

    try:
        with get_db() as conn:
            student = conn.execute(
                "SELECT name FROM students WHERE id = ?",
                (student_id,)
            ).fetchone()

            if student is None:
                flash("Student was not found.", "error")
                return redirect(url_for("dashboard"))

            conn.execute(
                "DELETE FROM attendance WHERE student_id = ?",
                (student_id,)
            )
            conn.execute(
                "DELETE FROM students WHERE id = ?",
                (student_id,)
            )

        try:
            face_path.unlink(missing_ok=True)
        except OSError as error:
            app.logger.warning(
                "Face photo cleanup failed for student %s: %s",
                student_id,
                error,
            )

        if FACE_MODEL_PATH.is_file():
            try:
                FACE_MODEL_PATH.unlink()
            except OSError as error:
                app.logger.warning("Face model cleanup failed: %s", error)
        FACE_MODEL_META_PATH.unlink(missing_ok=True)

        flash(
            f"{student['name']} and their attendance were deleted.",
            "success"
        )
    except sqlite3.Error:
        app.logger.exception("Student deletion failed for id %s", student_id)
        flash("The student could not be deleted. Try again.", "error")

    return redirect(url_for("dashboard"))


# ==================================================
# 8. SAVE STUDENT FACE PHOTO
# Supports both form data and JSON requests.
# ==================================================

@app.route("/students/capture-face", methods=["POST"])
@login_required
def capture_face():

    data = request.get_json(silent=True) or {}

    roll_number = str(
        data.get(
            "roll_number",
            request.form.get("roll_number", "")
        )
    ).strip()

    image_data = data.get(
        "face_image",
        request.form.get("face_image", "")
    )

    if not roll_number or not image_data:
        message = "Roll number and photo are both required."

        if request.is_json:
            return jsonify(success=False, message=message), 400

        flash(message, "error")
        return redirect(url_for("dashboard"))

    with get_db() as conn:
        student = conn.execute(
            "SELECT id, name FROM students WHERE roll_number = ?",
            (roll_number,)
        ).fetchone()

    if student is None:
        message = "Student was not found in the database."

        if request.is_json:
            return jsonify(success=False, message=message), 404

        flash(message, "error")
        return redirect(url_for("dashboard"))

    match = re.fullmatch(
        r"data:image/jpeg;base64,([A-Za-z0-9+/]*={0,2})",
        image_data
    )

    if not match:
        message = "Invalid photo format. Capture the photo again."

        if request.is_json:
            return jsonify(success=False, message=message), 400

        flash(message, "error")
        return redirect(url_for("dashboard"))

    try:
        image_bytes = base64.b64decode(
            match.group(1),
            validate=True
        )

        if not image_bytes or len(image_bytes) > 5 * 1024 * 1024:
            raise ValueError("Invalid image size")

        image = cv2.imdecode(
            np.frombuffer(image_bytes, dtype=np.uint8),
            cv2.IMREAD_COLOR
        )

        if image is None:
            raise ValueError("Invalid image")

        height, width = image.shape[:2]

        if width < 100 or height < 100:
            raise ValueError("Image resolution is too small")

        if get_face_crop(image) is None:
            raise ValueError("No face detected in photo")

        face_path = FACE_DIR / f"{student['id']}.jpg"

        if not cv2.imwrite(str(face_path), image):
            raise OSError("Image could not be saved")
        if FACE_MODEL_PATH.is_file():
            FACE_MODEL_PATH.unlink()
        FACE_MODEL_META_PATH.unlink(missing_ok=True)

        message = (
            f"Face photo saved for {student['name']}. "
            "Train the face model before recognition."
        )

        if request.is_json:
            return jsonify(success=True, message=message)

        flash(message, "success")
        return redirect(url_for("dashboard"))

    except (ValueError, OSError, cv2.error):
        message = "The photo could not be processed. Capture it again."

        if request.is_json:
            return jsonify(success=False, message=message), 400

        flash(message, "error")
        return redirect(url_for("dashboard"))


# ==================================================
# 9. FACE DETECTION HELPER
# ==================================================

def get_face_crop(image):

    if face_detector.empty():
        raise RuntimeError("The face detector could not be loaded.")

    gray = cv2.cvtColor(image, cv2.COLOR_BGR2GRAY)
    detection_gray = cv2.equalizeHist(gray)

    faces = face_detector.detectMultiScale(
        detection_gray,
        scaleFactor=1.08,
        minNeighbors=4,
        minSize=(70, 70)
    )

    if len(faces) == 0:
        return None

    x, y, w, h = max(
        faces,
        key=lambda box: box[2] * box[3]
    )

    padding_x = int(w * 0.18)
    padding_y = int(h * 0.22)
    left = max(0, x - padding_x)
    top = max(0, y - padding_y)
    right = min(gray.shape[1], x + w + padding_x)
    bottom = min(gray.shape[0], y + h + padding_y)
    face = gray[top:bottom, left:right]
    face = cv2.resize(face, (240, 240), interpolation=cv2.INTER_CUBIC)
    face = cv2.createCLAHE(
        clipLimit=2.0,
        tileGridSize=(8, 8),
    ).apply(face)

    return face


def get_face_variants(face):
    """Return controlled lighting and mirror variants for LBPH matching."""
    variants = [face, cv2.flip(face, 1)]
    gamma_lighter = np.power(face / 255.0, 0.8) * 255
    gamma_darker = np.power(face / 255.0, 1.2) * 255
    variants.extend([
        gamma_lighter.astype(np.uint8),
        gamma_darker.astype(np.uint8),
    ])
    return variants


# ==================================================
# 10. TRAIN FACE RECOGNITION MODEL
# ==================================================

@app.route("/attendance/train", methods=["POST"])
@login_required
def train_face_model():

    images = []
    labels = []

    with get_db() as conn:
        students = conn.execute(
            "SELECT id, name FROM students ORDER BY id"
        ).fetchall()

    for student in students:

        image_path = FACE_DIR / f"{student['id']}.jpg"

        if not image_path.is_file():
            continue

        image = cv2.imread(str(image_path))

        if image is None:
            continue

        face = get_face_crop(image)

        if face is not None:
            label = int(student["id"])
            variants = get_face_variants(face)
            images.extend(variants)
            labels.extend([label] * len(variants))

    if not images:
        flash(
            "No valid face photos were found. Capture clear, front-facing photos first.",
            "error"
        )
        return redirect(url_for("dashboard"))

    if not hasattr(cv2, "face"):
        flash(
            "OpenCV face recognition is unavailable. Install opencv-contrib-python.",
            "error",
        )
        return redirect(url_for("dashboard"))

    recognizer = cv2.face.LBPHFaceRecognizer_create(
        radius=2,
        neighbors=16,
        grid_x=8,
        grid_y=8,
    )

    recognizer.train(
        images,
        np.array(labels, dtype=np.int32)
    )

    recognizer.save(str(FACE_MODEL_PATH))
    FACE_MODEL_META_PATH.write_text(FACE_MODEL_VERSION, encoding="ascii")

    flash(
        f"Face model training complete. Training samples: {len(images)}",
        "success"
    )

    return redirect(url_for("dashboard"))


# ==================================================
# 11. RECOGNIZE FACE AND MARK ATTENDANCE
# ==================================================

@app.route("/attendance/recognize", methods=["POST"])
@login_required
def recognize_attendance():

    if not face_model_is_current():
        return jsonify({
            "success": False,
            "message": (
                "The face model is missing or out of date. "
                "Capture all photos, then click Train Face Model."
            )
        }), 400

    data = request.get_json(silent=True) or {}
    image_data = data.get("face_image", "")

    match = re.fullmatch(
        r"data:image/jpeg;base64,([A-Za-z0-9+/]*={0,2})",
        image_data
    )

    if not match:
        return jsonify({
            "success": False,
            "message": "A valid camera photo was not received."
        }), 400

    try:
        image_bytes = base64.b64decode(
            match.group(1),
            validate=True
        )

        if not image_bytes or len(image_bytes) > 5 * 1024 * 1024:
            raise ValueError("Invalid image size")

        image = cv2.imdecode(
            np.frombuffer(image_bytes, dtype=np.uint8),
            cv2.IMREAD_COLOR
        )

        if image is None:
            raise ValueError("Invalid image")

        face = get_face_crop(image)

        if face is None:
            return jsonify({
                "success": False,
                "message": "No face detected. Look directly at the camera."
            }), 400

        if not hasattr(cv2, "face"):
            raise RuntimeError(
                "OpenCV face recognition is unavailable."
            )

        recognizer = cv2.face.LBPHFaceRecognizer_create(
            radius=2,
            neighbors=16,
            grid_x=8,
            grid_y=8,
        )
        recognizer.read(str(FACE_MODEL_PATH))

        predictions = [
            recognizer.predict(variant)
            for variant in get_face_variants(face)
        ]
        best_prediction = min(predictions, key=lambda result: result[1])
        predicted_id, confidence = best_prediction
        matching_predictions = [
            result[1]
            for result in predictions
            if result[0] == predicted_id
        ]
        confidence = float(np.mean(matching_predictions))

        # Initial threshold only; real-world accuracy must be tested.
        if confidence > FACE_CONFIDENCE_THRESHOLD:
            return jsonify({
                "success": False,
                "message": (
                    "Face match was not confident enough. "
                    "Improve lighting and look at the camera."
                ),
                "confidence": round(float(confidence), 2),
            }), 422

        with get_db() as conn:

            student = conn.execute(
                """
                SELECT id, name, roll_number
                FROM students
                WHERE id = ?
                """,
                (int(predicted_id),)
            ).fetchone()

            if student is None:
                return jsonify({
                    "success": False,
                    "message": "The matching student was not found in the database."
                }), 404

            now = datetime.now()
            today = now.date().isoformat()
            current_time = now.time().replace(microsecond=0)
            current_time_text = current_time.isoformat()
            existing = conn.execute(
                """
                SELECT id, entry_time, exit_time
                FROM attendance
                WHERE student_id = ? AND attendance_date = ?
                """,
                (student["id"], today),
            ).fetchone()

            if existing is None:
                entry_status = (
                    "Early"
                    if current_time < COLLEGE_START
                    else "Outside Hours"
                    if current_time > COLLEGE_END
                    else "Late"
                    if current_time >= LATE_AFTER
                    else "Present"
                )
                conn.execute(
                    """
                    INSERT INTO attendance
                        (student_id, attendance_date, attendance_time,
                         entry_time, status)
                    VALUES (?, ?, ?, ?, ?)
                    """,
                    (
                        student["id"],
                        today,
                        current_time_text,
                        current_time_text,
                        entry_status,
                    ),
                )
                scan_type = "entry"
                message = (
                    f"Entry recorded for {student['name']} "
                    f"at {current_time.strftime('%I:%M %p')}."
                )
            elif existing["exit_time"] is None:
                completion_status = (
                    "Late / Completed"
                    if existing["entry_time"]
                    and datetime.strptime(
                        existing["entry_time"],
                        "%H:%M:%S",
                    ).time() >= LATE_AFTER
                    else "Completed"
                )
                conn.execute(
                    """
                    UPDATE attendance
                    SET exit_time = ?, status = ?
                    WHERE id = ?
                    """,
                    (current_time_text, completion_status, existing["id"]),
                )
                scan_type = "exit"
                message = (
                    f"Exit recorded for {student['name']} "
                    f"at {current_time.strftime('%I:%M %p')}."
                )
            else:
                scan_type = "duplicate"
                message = (
                    f"Entry and exit are already recorded for "
                    f"{student['name']} today."
                )

        return jsonify({
            "success": True,
            "message": message,
            "name": student["name"],
            "roll_number": student["roll_number"],
            "confidence": round(float(confidence), 2),
            "scan_type": scan_type,
            "already_marked": scan_type == "duplicate",
        })

    except (ValueError, OSError, cv2.error, RuntimeError) as error:

        app.logger.warning("Face recognition failed: %s", error)

        return jsonify({
            "success": False,
            "message": (
                "Recognition failed. Check lighting, camera position, "
                "face photo quality, and retrain the model."
            )
        }), 400


# ==================================================
# 12. ATTENDANCE RECORDS PAGE
# ==================================================

@app.route("/attendance/records")
@login_required
def attendance_records():

    with get_db() as conn:
        records = conn.execute(
            """
            SELECT
                students.name,
                students.roll_number,
                students.branch,
                students.academic_year,
                attendance.attendance_date,
                attendance.entry_time,
                attendance.exit_time,
                attendance.status
            FROM attendance
            JOIN students
                ON students.id = attendance.student_id
            ORDER BY attendance.id DESC
            """
        ).fetchall()

    return render_template(
        "attendance.html",
        records=records,
        username=session["admin"]
    )


# ==================================================
# 13. EXPORT ATTENDANCE TO EXCEL
# ==================================================

@app.route("/attendance/export.xlsx")
@login_required
def export_attendance_excel():
    """Export a populated attendance workbook with summary and roster sheets."""
    with get_db() as conn:
        records = conn.execute(
            """
            SELECT
                students.name AS student_name,
                students.roll_number AS roll_number,
                students.email AS email,
                students.branch AS branch,
                students.academic_year AS academic_year,
                attendance.attendance_date AS attendance_date,
                attendance.entry_time AS entry_time,
                attendance.exit_time AS exit_time,
                COALESCE(attendance.status, 'Not Marked') AS status
            FROM students
            LEFT JOIN attendance
                ON students.id = attendance.student_id
            ORDER BY attendance.attendance_date DESC,
                     attendance.attendance_time DESC,
                     LOWER(students.name) ASC
            """
        ).fetchall()
        attendance_days = conn.execute(
            "SELECT COUNT(*) AS total FROM attendance"
        ).fetchone()["total"]
        students = conn.execute(
            """
            SELECT name, roll_number, email, branch, academic_year, created_at
            FROM students
            ORDER BY LOWER(name)
            """
        ).fetchall()

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Attendance"
    headers = [
        "Student Name", "Roll Number", "Email", "Branch", "Academic Year",
        "Date", "Entry Time", "Exit Time", "Status",
    ]
    sheet.append(headers)

    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.alignment = Alignment(horizontal="center")

    for record in records:
        sheet.append([
            record["student_name"],
            record["roll_number"],
            record["email"] or "",
            record["branch"] or "",
            record["academic_year"] or "",
            record["attendance_date"],
            record["entry_time"] or "",
            record["exit_time"] or "",
            record["status"],
        ])

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column_cells in sheet.columns:
        max_length = max(len(str(cell.value or "")) for cell in column_cells)
        sheet.column_dimensions[column_cells[0].column_letter].width = min(max_length + 2, 35)

    roster = workbook.create_sheet("Student Roster")
    roster_headers = [
        "Student Name", "Roll Number", "Email", "Branch",
        "Academic Year", "Registered On",
    ]
    roster.append(roster_headers)
    for student in students:
        roster.append([
            student["name"],
            student["roll_number"],
            student["email"] or "",
            student["branch"] or "",
            student["academic_year"] or "",
            student["created_at"] or "",
        ])
    roster.freeze_panes = "A2"
    roster.auto_filter.ref = roster.dimensions
    for column_cells in roster.columns:
        max_length = max(len(str(cell.value or "")) for cell in column_cells)
        roster.column_dimensions[column_cells[0].column_letter].width = min(
            max_length + 2,
            35,
        )

    summary = workbook.create_sheet("Summary")
    summary.append(["AttendVision Attendance Report"])
    summary.append(["Generated", datetime.now().strftime("%Y-%m-%d %I:%M %p")])
    summary.append(["College Hours", "08:30 AM - 04:30 PM"])
    summary.append(["Registered Students", len(students)])
    summary.append(["Attendance Records", attendance_days])
    summary.append([])
    summary.append([
        "All registered students are listed in Attendance; "
        "unscanned students show Not Marked."
    ])
    summary.column_dimensions["A"].width = 32
    summary.column_dimensions["B"].width = 28

    workbook.active = 0

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    filename = f"AttendVision_Attendance_{datetime.now():%Y%m%d_%H%M%S}.xlsx"

    return send_file(
        output,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


# ==================================================
# 14. LOGOUT
# ==================================================

@app.route("/logout", methods=["POST"])
@login_required
def logout():

    session.clear()

    return redirect(url_for("login"))


# ==================================================
# 15. START APPLICATION
# ==================================================

if __name__ == "__main__":
    app.run(
        debug=os.environ.get("FLASK_DEBUG", "").lower() == "true"
    )
